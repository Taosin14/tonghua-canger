# -*- coding: utf-8 -*-
"""白族绘本素材模板.xlsx -> backend/db/02_seed_materials.sql

按素材 name 幂等(ON DUPLICATE KEY UPDATE),队友扩充到 60 条后重跑即可增量更新。
用法: python scripts/import_xlsx.py
"""
import json
import os
import re

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, '白族绘本素材模板 (3).xlsx')
OUT = os.path.join(ROOT, 'backend', 'db', '02_seed_materials.sql')

TYPE_SET = {'民间故事', '非遗工艺', '建筑场景', '人物角色', '节日民俗', '童谣儿歌'}


def esc(s):
    """SQL 单引号字符串转义(同时处理 JSON 字符串中的反斜杠)"""
    return (s or '').replace('\\', '\\\\').replace("'", "\\'")


def parse_age(cell):
    t = str(cell)
    m = re.search(r'(\d)\s*-\s*(\d)\s*岁', t)
    if m:
        return f'{m.group(1)}-{m.group(2)}'
    if '小班' in t:
        return '3-4'
    if '中班' in t:
        return '4-5'
    if '大班' in t:
        return '5-6'
    return '4-5'


def parse_goals(cell):
    """教育目标列 -> [{"domain","goal"}]"""
    text = str(cell or '').strip()
    if not text:
        return []
    goals = []
    for part in re.split(r'[;；\n]', text):
        part = part.strip()
        if not part:
            continue
        m = re.match(r'^(.*?)[:：](.*)$', part)
        if m:
            goals.append({'domain': m.group(1).strip(), 'goal': m.group(2).strip()})
        else:
            goals.append({'domain': '', 'goal': part})
    return goals


def parse_tags(cell):
    text = str(cell or '').strip()
    if not text:
        return []
    return [t.strip() for t in re.split(r'[\s,，]+', text) if t.strip()]


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    sheet = None
    for sn in wb.sheetnames:
        if '素材库' in sn and wb[sn].max_column >= 7:
            sheet = wb[sn]
            break
    if sheet is None:
        raise SystemExit('未找到素材库 sheet')

    lines = [
        '-- 素材种子数据(由 scripts/import_xlsx.py 生成,勿手改;重跑可增量更新)',
        'USE canger;',
        'SET NAMES utf8mb4;',
    ]
    count = 0
    last_mtype = ''
    for row in sheet.iter_rows(min_row=2, values_only=True):
        cells = [str(c or '').strip() if c is not None else '' for c in row[:7]]
        if len(cells) < 7:
            cells += [''] * (7 - len(cells))
        mtype, name, age, desc, prompt, goals, tags = cells
        if not mtype and last_mtype:
            mtype = last_mtype  # 素材类型为合并单元格, 向下填充
        if mtype in TYPE_SET:
            last_mtype = mtype
        if not name:
            continue
        if mtype not in TYPE_SET:
            print(f'[warn] 未知素材类型 "{mtype}" 于 "{name}",已跳过')
            continue
        v = {
            'type': esc(mtype),
            'name': esc(name),
            'age': esc(parse_age(age)),
            'desc': esc(desc),
            'prompt': esc(prompt),
            'goals': esc(json.dumps(parse_goals(goals), ensure_ascii=False)),
            'tags': esc(json.dumps(parse_tags(tags), ensure_ascii=False)),
        }
        lines.append(
            'INSERT INTO materials '
            '(material_type,name,age_band,description,image_prompt,edu_goals,tags,status) VALUES '
            f"('{v['type']}','{v['name']}','{v['age']}','{v['desc']}','{v['prompt']}',"
            f"'{v['goals']}','{v['tags']}','published') "
            'ON DUPLICATE KEY UPDATE '
            'material_type=VALUES(material_type), age_band=VALUES(age_band), '
            'description=VALUES(description), image_prompt=VALUES(image_prompt), '
            'edu_goals=VALUES(edu_goals), tags=VALUES(tags);'
        )
        count += 1

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines) + '\n')
    print(f'OK: {count} 条素材 -> {OUT}')


if __name__ == '__main__':
    main()
